# -*- coding: utf-8 -*-
import os, sys, docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import shutil

sys.stdout.reconfigure(encoding='utf-8')

md_text = """# DB INSURANCE TELEPHONE CALL SCRIPT & VERIFICATION PROTOCOL
## DB손해보험 유선 접수 및 자동차보험 가입증명서(서정대학교 제출용) 발급 요청 표준 통화 대본

---

**Document Control Number:** SCRIPT-DB-2026-0908-01  
**Target Subject:** Late MAHADI MONOAR HOSSAIN (마하디 모노아르 후세인 / Alien Registration No: `010530-7660021`)  
**Target Vehicle:** `28모1390` (Kia Morning / 기아 모닝 경승용)  
**Target Insurer:** DB Insurance Co., Ltd. (DB손해보험주식회사 / Customer Center: `1588-0100` / Overseas: `+82-2-2262-1234`)  
**Recipient Institution for Certificate:** Seojeong University International Affairs Office (서정대학교 국제교류처)  
**Execution Role:** Authorized Attorney-in-Fact / Bereaved Family Representative (선임 대리인 겸 유족 대표)  

---

# SECTION 1: CALL OVERVIEW & KEY IDENTIFIERS (통화 전 필수 확인 정보)

```
====================================================================================================
                                  CASE & POLICY IDENTIFIERS SHEET
====================================================================================================
1. Insured Name (피보험자 성명)     : MAHADI MONOAR HOSSAIN (마하디 모노아르 후세인)
2. Alien Registration No (외국인번호): 010530-7660021 (Date of Birth: May 30, 2001)
3. Subject Vehicle Plate (차량번호) : 28모1390 (Kia Morning / 기아 모닝 경승용)
4. Policy Renewal Date (계약갱신일) : January 21, 2026 (2026년 1월 21일 24:00 체결)
5. Date of Incident (사고일시)       : March 14, 2026 at 05:53 AM (2026년 3월 14일 05시 53분경)
6. Place of Incident (사고장소)     : Bukbun-ri 4-1, Hyeonnam-myeon, Yangyang-gun, Gangwon-do (7번 국도)
7. Academic Enrollment (학적사항)   : Seojeong University Smart Automobile Dept 1st Year (학번: 202581042)
8. Purpose of Request (발급목적)    : Submission to Seojeong University for D-2 Group Insurance & Record Clearance
====================================================================================================
```

---

# SECTION 2: TELEPHONE DIALING & ARS KEYPAD NAVIGATION

### Dialing Numbers:
* **Domestic (From within Korea):** `1588-0100` or `1600-0100`
* **International (From Overseas / Bangladesh):** `+82-2-2262-1234` or `+82-2-3011-5000`

### ARS Keypad Navigation Sequence:
1. When connected, listen to the automated prompt.
2. Press **`1`** (자동차 사고접수 및 보상문의 / Auto Claims & Inquiries).
3. Press **`9`** (상담원 연결 / Speak Directly with a Live Representative).

---

# SECTION 3: COMPLETE BILINGUAL CALL SCRIPT (한·영 병렬 통화 대본)

### Step 1: Opening & Situation Explanation (도입 및 상황 설명)

> **Caller (발신자):**  
> "안녕하세요. 지난 3월 14일 불의의 교통사고로 사망한 외국인 유학생 **마하디 모노아르 후세인 (MAHADI MONOAR HOSSAIN)**의 유족 대리인입니다.  
> 고인의 대학교인 **서정대학교(국제교류처)**에 사고 당시 자동차보험 가입 사실을 공식 증빙하는 서류를 제출해야 해서 연락드렸습니다.  
> 고인 소유 차량의 **'자동차보험 가입증명서(관공서/학교 제출용)'** 및 **'보험료 납입영수증'** 발급을 요청드립니다."

* **English Meaning:**  
  "Hello. I am calling on behalf of the family of foreign student MAHADI MONOAR HOSSAIN, who tragically passed away in a car accident on March 14. We need to submit an official document to his university (Seojeong University Office of International Affairs) confirming that he had an active auto insurance policy at the time of the accident. We request the issuance of the 'Certificate of Automobile Insurance' and 'Premium Payment Receipt'."

---

### Step 2: Providing Identification Details (인적사항 및 차량정보 제공)

> **DB Agent (상담원):**  
> *"네, 고인의 인적사항과 차량번호를 말씀해 주시겠습니까?"*  
> *(Yes, could you please provide the deceased's personal details and vehicle plate number?)*

> **Caller (발신자):**  
> "네, 불러드리겠습니다.  
> 1. **피보험자 성명:** 마하디 모노아르 후세인 (영문: MAHADI MONOAR HOSSAIN)  
> 2. **외국인등록번호:** `010530-7660021` (생년월일: 2001년 5월 30일)  
> 3. **차량번호:** **`28모1390`** (차종: 기아 모닝)  
> 4. **보험 갱신일자:** 2026년 1월 21일 갱신 체결 건입니다."

---

### Step 3: Making the Specific Document Request (서류 발급 방식 요청)

> **Caller (발신자):**  
> "고인이 사망하여 본인이 직접 웹사이트에서 본인인증(PASS/SMS)을 진행할 수 없는 상황입니다.  
> 서정대학교 외국인 유학생 단체보험 청구 및 학적 사망 처리를 위해, **2026년 3월 14일 사고 당시 유효하게 보험이 가입되어 있었다는 내용이 명시된 '자동차보험 가입증명서'**를 이메일로 발송해 주실 수 있습니까?"

* **English Meaning:**  
  "Because the insured has passed away, mobile 2FA self-verification on the website is impossible. To process Seojeong University's group accident insurance claim and student status clearance, can you please send the 'Certificate of Automobile Insurance' confirming active coverage as of the March 14, 2026 crash date to our email?"

---

### Step 4: Providing Email Coordinates & Collecting File Numbers (수신처 및 사건번호 확인)

> **DB Agent (상담원):**  
> *"발송받으실 이메일 주소나 팩스번호를 알려주시겠습니까?"*  
> *(Could you please provide the email address or fax number to receive the document?)*

> **Caller (발신자):**  
> "네, 서류를 수신할 공식 이메일 주소는 **`munimm247@gmail.com`** 입니다.  
> (m-u-n-i-m-m-2-4-7 @ gmail . com / 알파벳 확인)  
> 카카오톡 및 국내 연락처는 **`010-7209-1620`** (해외 직통: `+880-1941-646278`) 입니다.  
> 그리고 해당 건의 **'증권번호'**와 배정된 **'보상담당자 연락처(성명·직통전화·이메일)'**도 메모할 수 있도록 함께 말씀해 주시겠습니까?"

---

### Step 5: Closing (통화 마감)

> **Caller (발신자):**  
> "친절하게 확인해 주셔서 대단히 감사합니다. 이메일로 서류가 도착하는 대로 바로 대학교에 제출하도록 하겠습니다. 수고하세요."

---

# SECTION 4: ANTICIPATED AGENT INQUIRIES & BULLETPROOF RESPONSES (예상 질문 및 답변)

```
┌─────────────────────────────────────────────────────────────────────────────────┐
│                     ANTICIPATED AGENT Q&A HANDLING GUIDE                        │
├─────────────────────────────────────────────────────────────────────────────────┤
│                                                                                 │
│  Q1. 고인과의 관계가 어떻게 되시나요? (What is your relationship to deceased?)   │
│  A1. "방글라데시에 계신 고인의 부모님(법정상속인)으로부터 정식 권한을 위임받은    │
│      선임 수임인 겸 유족 지원 대리인입니다."                                     │
│                                                                                 │
│  Q2. 사망 사실이나 위임 서류를 보내주실 수 있나요? (Can you provide proof?)     │
│  A2. "네, 강릉아산병원 사망진단서, 속초경찰서 사체인도지휘서 및 유족 특별위임장을  │
│      모두 구비하고 있으며, 요청하시는 팩스나 이메일로 즉시 송부해 드리겠습니다."│
│                                                                                 │
│  Q3. 대학교에서 이 서류를 왜 요구하나요? (Why does the university need this?)   │
│  A3. "서정대학교 외국인 유학생 단체상해보험 청구 및 유학생 사망 사고 처리를 위해 │
│      고인이 개인 자동차보험에 가입되어 있었는지 공식 확인이 필요하기 때문입니다."│
│                                                                                 │
└─────────────────────────────────────────────────────────────────────────────────┘
```

---

# SECTION 5: CONTACT DIRECTORY & ESCALATION DESK

| Department / Agency | Contact Number | Primary Purpose |
|---|---|---|
| **DB Insurance Overseas Customer Desk** | `+82-2-2262-1234` | Direct overseas intake for policy certificate requests |
| **DB Auto Claims & Emergency Hotline** | `+82-2-3011-5000` / `1588-0100` | Auto accident claim reporting & adjuster assignment |
| **DB Central Claims Document Fax** | `0505-181-4800` / `0505-181-4810` | Direct fax transmission for POA and death certificates |
| **Seojeong University International Office** | `+82-31-860-5000` (`intl@seojeong.ac.kr`) | Recipient institution for group insurance & student file |
| **Gangwon Sokcho Police Station Traffic Team**| `+82-33-634-0357` (Case `2026-00020`) | Police investigation confirmation & accident report |

---

*Document prepared for Fiduciary Attorney-in-Fact HASIBUL MUNIM for cross-border estate recovery.*
"""

md_path = r'd:\promotion\legal_claim_package\14_db_insurance_telephone_call_script_for_policy_certificate.md'
docx_path = r'd:\promotion\legal_claim_package\docx\14_db_insurance_telephone_call_script_for_policy_certificate.docx'
pdf_path = r'd:\promotion\legal_claim_package\pdf\14_db_insurance_telephone_call_script_for_policy_certificate.pdf'

with open(md_path, 'w', encoding='utf-8') as f:
    f.write(md_text)
print('Saved Markdown:', md_path)

# Build DOCX
doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

lines = md_text.split('\n')
for line in lines:
    if line.startswith('# '):
        p = doc.add_heading(line[2:], level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif line.startswith('## '):
        p = doc.add_heading(line[3:], level=2)
    elif line.startswith('### '):
        p = doc.add_heading(line[4:], level=3)
    elif line.startswith('> '):
        p = doc.add_paragraph()
        run = p.add_run(line[2:])
        run.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
        p.paragraph_format.left_indent = Inches(0.3)
    elif line.startswith('* ') or line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(line[2:])
    elif line.startswith('---'):
        p = doc.add_paragraph()
        p.add_run('_________________________________________________________________________________')
    else:
        if line.strip():
            doc.add_paragraph(line)

doc.save(docx_path)
print('Saved DOCX:', docx_path)

# Convert to PDF via Word COM
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    doc_com = word.Documents.Open(os.path.abspath(docx_path))
    doc_com.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
    doc_com.Close()
    print('Exported PDF:', pdf_path)
finally:
    word.Quit()

# Copy to 03 pack
pack_03_pdf = r'd:\promotion\03_KOREAN_LAWYERS_AND_INSURERS_FILING_PACK\14_DB_Insurance_Telephone_Call_Script_for_Policy_Certificate.pdf'
shutil.copy2(pdf_path, pack_03_pdf)
print('Copied to Pack 03:', pack_03_pdf)
