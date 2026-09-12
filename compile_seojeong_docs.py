# -*- coding: utf-8 -*-
import os, sys, docx, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

def convert_md_to_docx_and_pdf(md_file_path, docx_file_path, pdf_file_path):
    with open(md_file_path, 'r', encoding='utf-8') as f:
        md_content = f.read()

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    lines = md_content.split('\n')
    for line in lines:
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('## '):
            p = doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            p = doc.add_heading(line[5:], level=4)
        elif line.startswith('> '):
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            run.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
            p.paragraph_format.left_indent = Inches(0.3)
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(line[2:])
        elif line.startswith('---'):
            p = doc.add_paragraph()
            p.add_run('_________________________________________________________________________________')
        else:
            if line.strip():
                doc.add_paragraph(line)

    doc.save(docx_file_path)
    print(f'Saved DOCX: {docx_file_path}')

    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc_com = word.Documents.Open(os.path.abspath(docx_file_path))
        doc_com.SaveAs(os.path.abspath(pdf_file_path), FileFormat=17)
        doc_com.Close()
        print(f'Exported PDF: {pdf_file_path}')
    finally:
        word.Quit()

if __name__ == '__main__':
    files = [
        ('15_seojeong_university_telephone_call_script', '15_Seojeong_University_Telephone_Call_Script.pdf'),
        ('16_seojeong_university_email_and_epeople_petition_package', '16_Seojeong_University_Email_and_ePeople_Petition_Package.pdf')
    ]
    for base_name, dest_pdf_name in files:
        md_p = f'd:/promotion/legal_claim_package/{base_name}.md'
        docx_p = f'd:/promotion/legal_claim_package/docx/{base_name}.docx'
        pdf_p = f'd:/promotion/legal_claim_package/pdf/{base_name}.pdf'
        convert_md_to_docx_and_pdf(md_p, docx_p, pdf_p)
        
        # Copy to pack 03
        dest_pack = f'd:/promotion/03_KOREAN_LAWYERS_AND_INSURERS_FILING_PACK/{dest_pdf_name}'
        shutil.copy2(pdf_p, dest_pack)
        print(f'Copied to pack 03: {dest_pack}')

    print('All Seojeong documents successfully compiled!')
