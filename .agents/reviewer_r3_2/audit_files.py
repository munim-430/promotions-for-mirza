import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import docx
import pypdf

base_dir = r'd:/promotion/legal_claim_package'
docx_dir = os.path.join(base_dir, 'docx')
pdf_dir = os.path.join(base_dir, 'pdf')

md_files = sorted([f for f in os.listdir(base_dir) if f.endswith('.md')])
docx_files = sorted([f for f in os.listdir(docx_dir) if f.endswith('.docx')])
pdf_files = sorted([f for f in os.listdir(pdf_dir) if f.endswith('.pdf')])

print('MD count:', len(md_files))
print('DOCX count:', len(docx_files))
print('PDF count:', len(pdf_files))
print('=' * 120)
header = '{:<55} | {:<9} | {:<25} | {:<20}'.format('File Name', 'MD Bytes', 'DOCX (p/t/chars)', 'PDF (pg/chars)')
print(header)
print('=' * 120)

for md_name in md_files:
    stem = os.path.splitext(md_name)[0]
    md_path = os.path.join(base_dir, md_name)
    docx_path = os.path.join(docx_dir, stem + '.docx')
    pdf_path = os.path.join(pdf_dir, stem + '.pdf')
    
    md_size = os.path.getsize(md_path)
    
    if os.path.exists(docx_path):
        doc = docx.Document(docx_path)
        doc_paras = len(doc.paragraphs)
        doc_tables = len(doc.tables)
        doc_chars = sum(len(p.text) for p in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    doc_chars += len(cell.text)
        docx_info = f'{doc_paras}p / {doc_tables}t / {doc_chars}c'
    else:
        docx_info = 'MISSING'
        
    if os.path.exists(pdf_path):
        reader = pypdf.PdfReader(pdf_path)
        pdf_pages = len(reader.pages)
        pdf_chars = sum(len(p.extract_text() or '') for p in reader.pages)
        pdf_info = f'{pdf_pages} pgs / {pdf_chars}c'
    else:
        pdf_info = 'MISSING'
        
    line = '{:<55} | {:<9} | {:<25} | {:<20}'.format(stem, md_size, docx_info, pdf_info)
    print(line)

print('=' * 120)
