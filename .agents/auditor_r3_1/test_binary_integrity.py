import os
import sys
import glob
import fitz # PyMuPDF
import docx

sys.stdout.reconfigure(encoding='utf-8')

PACKAGE_DIR = r"d:/promotion/legal_claim_package"
DOCX_DIR = os.path.join(PACKAGE_DIR, "docx")
PDF_DIR = os.path.join(PACKAGE_DIR, "pdf")

print("=== DETAILED DOCX AUDIT ===")
docx_files = sorted(glob.glob(os.path.join(DOCX_DIR, "*.docx")))
for df in docx_files:
    fname = os.path.basename(df)
    doc = docx.Document(df)
    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    total_text_len = sum(len(p.text) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                total_text_len += len(cell.text)
    print(f"[DOCX] {fname:60s} | Size: {os.path.getsize(df):7,d} bytes | Paras: {para_count:4d} | Tables: {table_count:2d} | Text Chars: {total_text_len:7,d}")

print("\n=== DETAILED PDF AUDIT ===")
pdf_files = sorted(glob.glob(os.path.join(PDF_DIR, "*.pdf")))
for pf in pdf_files:
    fname = os.path.basename(pf)
    doc = fitz.open(pf)
    page_count = len(doc)
    total_chars = 0
    has_images = False
    for page in doc:
        text = page.get_text()
        total_chars += len(text)
        if page.get_images():
            has_images = True
    print(f"[PDF]  {fname:60s} | Size: {os.path.getsize(pf):7,d} bytes | Pages: {page_count:3d} | Text Chars: {total_chars:7,d}")
