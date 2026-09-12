# ==============================================================================
# SCRIPT: compile_doc.py
# PURPOSE: Compiles 11_dhaka_to_seoul_dhl_dispatch_and_10day_tracking_guide.md to DOCX and PDF
# ==============================================================================

import os
import re
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

sys.stdout.reconfigure(encoding='utf-8')

COLOR_PRIMARY = RGBColor(0, 32, 96)       # Deep Navy
COLOR_SECONDARY = RGBColor(80, 80, 80)    # Charcoal Gray
COLOR_BODY = RGBColor(30, 30, 30)         # Off Black
COLOR_CALLOUT_BG = "F2F4F7"               # Soft Slate

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_styled_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Legal Claim Recovery Dossier | Case No. 2026-SF-AUTO-28M1390")
        f_run.font.name = "Malgun Gothic"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(128, 128, 128)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Malgun Gothic'
    style_normal.font.size = Pt(10.0)
    style_normal.font.color.rgb = COLOR_BODY
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(3)
    
    return doc

def parse_markdown_to_docx(md_path, output_docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    doc = create_styled_document()
    
    in_code_block = False
    code_buffer = []
    in_table = False
    table_buffer = []
    
    for line in lines:
        raw_line = line.rstrip('\r\n')
        
        if raw_line.startswith('```'):
            if in_code_block:
                in_code_block = False
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, COLOR_CALLOUT_BG)
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.line_spacing = 1.15
                crun = cp.add_run('\n'.join(code_buffer))
                crun.font.name = 'Consolas'
                crun.font.size = Pt(8.0)
                crun.font.color.rgb = RGBColor(40, 40, 40)
                
                doc.add_paragraph()
                code_buffer = []
            else:
                in_code_block = True
                code_buffer = []
            continue
            
        if in_code_block:
            code_buffer.append(raw_line)
            continue
            
        if '|' in raw_line and not raw_line.startswith('#'):
            if not in_table:
                in_table = True
                table_buffer = []
            table_buffer.append(raw_line)
            continue
        else:
            if in_table:
                in_table = False
                rows_data = []
                for tline in table_buffer:
                    if re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', tline):
                        continue
                    cells = [c.strip() for c in tline.strip('|').split('|')]
                    if cells:
                        rows_data.append(cells)
                
                if rows_data:
                    num_cols = max(len(r) for r in rows_data)
                    table = doc.add_table(rows=len(rows_data), cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    for r_idx, row in enumerate(rows_data):
                        for c_idx, val in enumerate(row):
                            if c_idx < num_cols:
                                c = table.cell(r_idx, c_idx)
                                c.text = val
                                cp = c.paragraphs[0]
                                cp.paragraph_format.space_after = Pt(2)
                                cp.paragraph_format.line_spacing = 1.1
                                crun = cp.runs[0]
                                crun.font.name = 'Malgun Gothic'
                                
                                if r_idx == 0:
                                    set_cell_background(c, "E8ECF2")
                                    crun.font.bold = True
                                    crun.font.size = Pt(9.0)
                                    crun.font.color.rgb = COLOR_PRIMARY
                                else:
                                    crun.font.size = Pt(8.5)
                                set_cell_margins(c, top=60, bottom=60, left=80, right=80)
                    doc.add_paragraph()
                table_buffer = []
                
        if not raw_line.strip():
            continue
            
        if raw_line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(raw_line[2:].strip())
            run.font.name = 'Malgun Gothic'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
        elif raw_line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(raw_line[3:].strip())
            run.font.name = 'Malgun Gothic'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
        elif raw_line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(raw_line[4:].strip())
            run.font.name = 'Malgun Gothic'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY
        elif raw_line.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('―' * 40)
            run.font.color.rgb = RGBColor(180, 180, 180)
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', raw_line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.name = 'Malgun Gothic'
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                    run.font.name = 'Malgun Gothic'

    doc.save(output_docx_path)
    print(f"[OK] Generated DOCX: {output_docx_path}")

def convert_docx_to_pdf_via_word(docx_path, pdf_path):
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        abs_docx = os.path.abspath(docx_path)
        abs_pdf = os.path.abspath(pdf_path)
        wb = word.Documents.Open(abs_docx)
        wb.SaveAs2(abs_pdf, FileFormat=17)
        wb.Close()
        print(f"[OK] Exported PDF: {abs_pdf}")
    finally:
        word.Quit()

def verify_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    print(f"[VERIFY] PDF '{os.path.basename(pdf_path)}': {len(doc)} pages, FileSize: {os.path.getsize(pdf_path)} bytes")
    doc.close()

if __name__ == '__main__':
    target_md = 'd:/promotion/legal_claim_package/11_dhaka_to_seoul_dhl_dispatch_and_10day_tracking_guide.md'
    out_docx = 'd:/promotion/legal_claim_package/docx/11_dhaka_to_seoul_dhl_dispatch_and_10day_tracking_guide.docx'
    out_pdf = 'd:/promotion/legal_claim_package/pdf/11_dhaka_to_seoul_dhl_dispatch_and_10day_tracking_guide.pdf'
    
    parse_markdown_to_docx(target_md, out_docx)
    convert_docx_to_pdf_via_word(out_docx, out_pdf)
    verify_pdf(out_pdf)
